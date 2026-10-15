"""
evals/build_corpus.py — Generate the synthetic evaluation corpus.

    python -m evals.build_corpus            # writes evals/corpus/ and evals/tabular/
    python -m evals.build_corpus --check    # regenerate in a temp dir and compare checksums

Every document is fictional ("Northwind Analytics" and partners) and written
here, so each benchmark label can be traced to an exact planted sentence. No
real people are named; approvers are identified by role.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from evals.pdfgen import PAGE_H, TextLine, flow, write_pdf

CORPUS_VERSION = "1.0.0"
ROOT = Path(__file__).resolve().parent
CORPUS_DIR = ROOT / "corpus"
TABULAR_DIR = ROOT / "tabular"

TRAVEL_2023 = """Northwind Analytics - Travel and Expense Policy (2023 edition)

Scope. This policy applies to all employees travelling on company business. It was approved by the Finance Committee and took effect on 1 January 2023.

Meals. Employees may claim a daily meal allowance of 45 dollars for domestic travel. For international travel the daily allowance is 75 dollars. Alcohol is not reimbursable.

Lodging. Hotel stays up to 200 dollars per night are approved automatically. Stays above 200 dollars per night require written approval from the employee's manager before booking.

Air travel. Economy class is required for all flights under six hours. Business class may be booked for flights of six hours or longer.

Ground transport. Personal car use is reimbursed at 58 cents per mile. Rental cars must be compact class unless three or more employees share the vehicle.

Submitting expenses. Expense reports must be submitted within 30 days of the end of the trip. Itemised receipts are required for any single expense above 25 dollars."""

TRAVEL_2024 = """Northwind Analytics - Travel and Expense Policy (2024 edition)

Scope. This edition replaces the 2023 policy from 1 March 2024. It applies to all employees travelling on company business.

Meals. The daily meal allowance for domestic travel is increased to 60 dollars. For international travel the daily allowance is 90 dollars. Alcohol remains non-reimbursable.

Lodging. Hotel stays up to 250 dollars per night are approved automatically. Stays above 250 dollars per night require approval from a director, not a manager.

Air travel. Economy class is required for flights under eight hours. Business class may be booked for flights of eight hours or longer.

Ground transport. Personal car use is reimbursed at 67 cents per mile. Rail travel is preferred over flying for trips under 200 miles.

Submitting expenses. Expense reports must be submitted within 15 days of the end of the trip. Itemised receipts are still required for any single expense above 25 dollars."""

REMOTE_ADDENDUM = """Northwind Analytics - Remote Work Addendum (2024)

Effective date. This addendum takes effect on 1 June 2024 and amends the remote work section of the Employee Handbook.

Eligibility. Employees who have completed their probation period may work remotely up to three days per week. Employees still in probation keep the Handbook limit.

Location. Remote work must be performed from within the employee's country of employment. Working from another country requires approval from the People Team at least 30 days in advance.

Equipment. Eligible remote employees receive a one-time home office stipend of 300 dollars. Company laptops used at home remain subject to the IT Security Policy."""

BENEFITS_FAQ = """Northwind Analytics - Benefits FAQ

When must I enrol in health insurance? New employees must enrol within 30 days of their start date. Missing the window means waiting for the next annual enrolment period.

How much of the premium does the company pay? The company pays 80 percent of the health insurance premium for employees and 50 percent for dependants.

Is dental cover included? Yes, dental cover is included in the standard health plan. Vision cover is optional and paid entirely by the employee.

What is the retirement plan match? The company matches retirement plan contributions up to 4 percent of base salary.

Is there a wellness benefit? Employees receive a wellness stipend of 50 dollars per month, paid with salary.

How much life insurance is provided? Life insurance cover equal to two times annual base salary is provided at no cost."""

HANDBOOK = [
    ("Employee Handbook - Northwind Analytics", None),
    ("Probation", "All new employees complete a probation period of 90 days. Probation reviews are held by the line manager before the period ends."),
    ("Working hours", "Core hours are 10:00 to 15:00, Monday to Friday. Outside core hours employees may arrange their schedule with their manager."),
    ("Annual leave", "Full-time employees accrue 1.5 days of annual leave per month, for a total of 18 days per year. Up to 5 unused days may be carried into the next year. Unused days above that are forfeited on 31 December."),
    ("Parental leave", "The primary caregiver receives 16 weeks of fully paid parental leave. The secondary caregiver receives 4 weeks of fully paid parental leave. Leave must start within 12 months of the birth or adoption."),
    ("Remote work", "Employees may work remotely up to two days per week with their manager's agreement."),
    ("Conduct", "Employees must complete the annual code of conduct training by 31 March each year."),
]

IT_SECURITY = """Northwind Analytics - IT Security Policy

Passwords. Passwords must be at least 14 characters long and must not reuse any of the previous five passwords. Passwords must be changed immediately if a compromise is suspected.

Multi-factor authentication. Multi-factor authentication is required for email, the VPN and all administrative consoles.

Devices. Company laptops must use full-disk encryption. Screens must lock automatically after 5 minutes of inactivity. USB storage devices are blocked by default; exceptions require approval from the security team.

Incidents. Lost or stolen devices must be reported to the IT service desk within 24 hours. Suspected phishing emails must be reported using the report button in the mail client.

Training. All employees must complete security awareness training once a year. New employees must complete it within their first 30 days."""

VENDOR_CONTRACT = """Master Services Agreement - Northwind Analytics and Brightline Logistics

Term. This agreement starts on 1 January 2024 and runs for an initial term of 24 months. It renews automatically for 12-month periods unless either party gives notice.

Payment. Invoices are payable within 45 days of receipt. Late payments accrue interest at 1.5 percent per month.

Termination. Either party may terminate the agreement for convenience with 60 days' written notice. Either party may terminate immediately for a material breach that is not cured within 15 days.

Liability. Each party's total liability under this agreement is capped at 500,000 dollars, except for breaches of confidentiality.

Data. Brightline Logistics must store Northwind shipment data in data centres located in the European Union.

Governing law. This agreement is governed by the laws of the State of Delaware."""

# Two-column handbook: the over-21 and under-21 rules sit side by side, as in the
# real document that produced the Phase 1B verifier false positive.
DRIVER_LEFT = """Drivers aged 21 and over. A driver aged 21 or over commits an offence if their blood alcohol concentration is 0.08 percent or higher. A first offence leads to a licence revocation of at least six months and a fine.

Seat belts. Seat belts must be worn by every occupant of a fleet vehicle at all times, including short trips between company sites.

Phones. Handheld phone use while driving a fleet vehicle is prohibited. Hands-free calls are allowed only when the vehicle is stationary."""

DRIVER_RIGHT = """Drivers under age 21. Drivers under age 21 are covered by the zero tolerance rule. A blood alcohol concentration from 0.02 to 0.07 percent leads to a six month licence suspension. A second violation leads to revocation for one year.

Incidents. Any collision involving a fleet vehicle must be reported to the fleet coordinator within 24 hours, even if no one is injured.

Lights. Headlights must be switched on whenever windscreen wipers are in use."""

DRIVER_TABLE = [
    ("Vehicle class", "Check", "Interval"),
    ("Van", "Brakes", "12 months"),
    ("Truck", "Tyres", "6 months"),
    ("Car", "Lights", "12 months"),
]

FLEET_CSV = """region,vehicle_type,miles,trips,fuel_litres
EU,van,1200,30,140
EU,truck,800,10,210
EU,car,650,25,48
US,van,1500,40,180
US,truck,2100,18,520
APAC,car,900,35,70
APAC,van,400,12,52
"""


def _write_text(path: Path, text: str) -> None:
    path.write_text(text.strip() + "\n", encoding="utf-8")


def _write_handbook(path: Path) -> None:
    from docx import Document

    doc = Document()
    for heading, body in HANDBOOK:
        if body is None:
            doc.add_heading(heading, level=1)
        else:
            doc.add_heading(heading, level=2)
            doc.add_paragraph(body)
    fixed = datetime(2024, 1, 1)
    props = doc.core_properties
    props.created = props.modified = props.last_printed = fixed
    props.author = props.last_modified_by = "Northwind Analytics"
    props.revision = 1
    doc.save(path)
    _normalize_zip(path)


def _normalize_zip(path: Path) -> None:
    """Rewrite a zip-based file (DOCX) with fixed entry timestamps for reproducible bytes."""
    import zipfile

    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info.filename)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(2024, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, data)


def _single_column_pdf(path: Path, text: str) -> None:
    lines = flow(text, 72, PAGE_H - 72, 92)
    pages, page = [], []
    offset = 0.0
    for ln in lines:
        y = ln.y + offset
        if y < 72:
            pages.append(page)
            page = []
            offset += (PAGE_H - 72) - y
            y = ln.y + offset
        page.append(TextLine(ln.x, y, ln.text, ln.size))
    pages.append(page)
    write_pdf(path, pages)


def _driver_pdf(path: Path) -> None:
    top = PAGE_H - 72
    lines = [TextLine(72, top, "Fleet Driver Safety Handbook - Rules for Company Vehicles", 14)]
    top -= 32
    lines += flow(DRIVER_LEFT, 54, top, 44) + flow(DRIVER_RIGHT, 318, top, 44)
    y = min(ln.y for ln in lines) - 36
    lines.append(TextLine(72, y, "Inspection schedule", 12))
    y -= 20
    for row in DRIVER_TABLE:
        lines += [TextLine(72, y, row[0]), TextLine(250, y, row[1]), TextLine(430, y, row[2])]
        y -= 14
    write_pdf(path, [lines])


def build(corpus_dir: Path = CORPUS_DIR, tabular_dir: Path = TABULAR_DIR) -> dict:
    corpus_dir.mkdir(parents=True, exist_ok=True)
    tabular_dir.mkdir(parents=True, exist_ok=True)
    _write_text(corpus_dir / "travel_policy_2023.txt", TRAVEL_2023)
    _write_text(corpus_dir / "travel_policy_2024.txt", TRAVEL_2024)
    _write_text(corpus_dir / "remote_work_addendum_2024.txt", REMOTE_ADDENDUM)
    _write_text(corpus_dir / "benefits_faq.txt", BENEFITS_FAQ)
    _write_handbook(corpus_dir / "employee_handbook.docx")
    _single_column_pdf(corpus_dir / "it_security_policy.pdf", IT_SECURITY)
    _single_column_pdf(corpus_dir / "vendor_contract_brightline.pdf", VENDOR_CONTRACT)
    _driver_pdf(corpus_dir / "driver_safety_handbook.pdf")
    (tabular_dir / "fleet_usage.csv").write_text(FLEET_CSV, encoding="utf-8")

    manifest = {"version": CORPUS_VERSION, "files": {}}
    for directory in (corpus_dir, tabular_dir):
        for f in sorted(directory.iterdir()):
            if f.name != "MANIFEST.json":
                manifest["files"][f"{directory.name}/{f.name}"] = hashlib.sha256(f.read_bytes()).hexdigest()
    (corpus_dir / "MANIFEST.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="rebuild in a temp dir and compare with the committed manifest")
    args = ap.parse_args()
    if not args.check:
        manifest = build()
        print(f"Wrote {len(manifest['files'])} files, corpus version {manifest['version']}")
        return 0
    committed = json.loads((CORPUS_DIR / "MANIFEST.json").read_text())
    with tempfile.TemporaryDirectory() as tmp:
        rebuilt = build(Path(tmp) / "corpus", Path(tmp) / "tabular")
    diffs = [k for k in committed["files"] if committed["files"][k] != rebuilt["files"].get(k)]
    print("identical rebuild" if not diffs else f"files differ on rebuild: {diffs}")
    return 1 if diffs else 0


if __name__ == "__main__":
    sys.exit(main())
